"""Convierte el catálogo de Z E-Scooters (xlsx) en public/data/catalogo.json.

Uso:  python scripts/catalogo_a_json.py "Catalogo_Z_eScooters_MX.xlsx"

Precios del Excel = precio público (IVA incluido). El buscador calcula el
precio sin IVA al agregar el concepto a la factura.
"""
import json, re, sys, unicodedata
from datetime import date
from pathlib import Path
import openpyxl

SECCIONES = {
    "Scooters": ("scooters", "Scooters"),
    "Monociclos EUC": ("monociclos", "Monociclos"),
    "Nosfet": ("nosfet", "Nosfet"),
    "Hyvel (Baterías y cargadores)": ("energia", "Baterías y cargadores"),
    "Baterías (lista anterior)": ("energia", "Baterías y cargadores"),
    "Refacciones y accesorios": ("refacciones", "Refacciones y accesorios"),
}
# Claves SAT (c_ClaveProdServ). Editables en la factura.
SAT_VEHICULO = "25101900"
SAT_BATERIA = "26111701"
SAT_REFACCION = "25174000"


def slug(text):
    t = unicodedata.normalize("NFD", text).encode("ascii", "ignore").decode().lower()
    return re.sub(r"[^a-z0-9]+", "-", t).strip("-")


def estado(raw):
    s = (raw or "").lower()
    if "agotado" in s: return "agotado"
    if "pre orden" in s or "preorden" in s: return "preorden"
    if "verificar" in s: return "verificar"
    if "oferta" in s: return "oferta"
    if "a partir" in s: return "desde"
    return "disponible"


MARCAS = {"KINGSONG": "KingSong", "KAABO": "Kaabo", "LANGFEITE": "Langfeite", "INMOTION": "Inmotion", "NOSFET": "Nosfet"}


def descripcion(cat, tipo, modelo):
    if cat in ("scooters", "nosfet", "monociclos"):
        clase = "Monociclo eléctrico" if cat == "monociclos" else "Scooter eléctrico"
        marca = MARCAS.get(tipo.upper(), tipo)
        if modelo.lower().startswith(marca.lower()):
            return f"{clase} {modelo}"
        return f"{clase} {marca} {modelo}"
    return modelo


def main(src, out):
    wb = openpyxl.load_workbook(src, data_only=True)
    items, vistos = [], set()
    for hoja, (cat, cat_label) in SECCIONES.items():
        if hoja not in wb.sheetnames: continue
        for tipo, modelo, precio, tprecio, url in wb[hoja].iter_rows(min_row=2, max_col=5, values_only=True):
            if not modelo or precio in (None, ""): continue
            tipo, modelo = str(tipo or "").strip(), str(modelo).strip()
            clave = (slug(modelo), round(float(precio), 2))
            if clave in vistos: continue
            vistos.add(clave)
            t = tipo.lower()
            if cat in ("scooters", "monociclos", "nosfet"):
                sat, marca, linea = SAT_VEHICULO, tipo.upper(), MARCAS.get(tipo.upper(), tipo)
            elif cat == "energia":
                es_bat = t.startswith(("batería", "kit", "pack")) and not modelo.lower().startswith(("caja", "cargador", "conector"))
                sat = SAT_BATERIA if es_bat else SAT_REFACCION
                marca = "HYVEL" if hoja.startswith("Hyvel") else "Z"
                linea = tipo
            else:
                generico = t in ("refacción", "accesorio")
                sat, marca, linea = SAT_REFACCION, ("" if generico else tipo.upper()), tipo
            items.append({
                "id": f"{cat}-{slug(modelo)}-{len(items)}",
                "cat": cat,
                "cat_label": cat_label,
                "brand": marca,
                "line": linea,
                "model": modelo,
                "description": descripcion(cat, tipo, modelo),
                "product_key": sat,
                "unit_key": "H87",
                "price": round(float(precio), 2),
                "price_label": str(tprecio or "").strip(),
                "status": estado(tprecio),
                "legacy": hoja == "Baterías (lista anterior)",
                "url": str(url or "").strip(),
            })
    data = {"source": Path(src).name, "generated": date.today().isoformat(), "tax_included": True, "count": len(items), "items": items}
    Path(out).write_text(json.dumps(data, ensure_ascii=False, separators=(",", ":")), encoding="utf-8")
    print(f"{len(items)} productos → {out}")


if __name__ == "__main__":
    main(sys.argv[1], sys.argv[2] if len(sys.argv) > 2 else "public/data/catalogo.json")
